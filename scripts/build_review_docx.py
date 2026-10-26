#!/usr/bin/env python3
"""Build the client-facing website copy review document.

Usage:
  build_review_docx.py --pages pages.json --items items.json --out Review.docx

pages.json : page order + metadata + cover "meta" block (see assets/pages.example.json)
items.json : output of merge_and_verify.py (items carry "ref")

After saving, the script re-opens the document and asserts that the number of data rows equals the
number of items and that the "Suggested revision" and "Notes" columns are empty in every data row.
"""
import argparse
import json
import re
import sys
from datetime import date

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_BREAK, WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

COLS = [("Ref", 2.2), ("Element", 4.3), ("Current copy", 8.4), ("Suggested revision", 8.4), ("Notes", 3.6)]
HEADER_FILL = "1F2937"
BANNER_FILL = "DCE6F0"
ZEBRA_FILL = "F5F7FA"
DYNAMIC_FILL = "FFF1CC"
INK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x5B, 0x65, 0x73)
ACCENT = RGBColor(0x1D, 0x4E, 0x89)
REF_RE = re.compile(r"^[A-Z0-9_]+-\d{2,}$")


# ---------- low-level helpers ----------

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:shd")):
        tcPr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))


def fixed_layout(table, widths_cm):
    tblPr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    grid = table._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths_cm):
        gc.set(qn("w:w"), str(int(Cm(w).twips)))
    table.autofit = False


def set_cell_margins(table, top=60, bottom=60, left=90, right=90):
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(v))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.append(mar)


def set_borders(table, color="C9D1DB", size=4):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tblPr.append(borders)


def write(cell_or_par, text, size=9, bold=False, italic=False, color=None, new_par=False):
    """Write text into a cell (first paragraph unless new_par) or a paragraph."""
    if hasattr(cell_or_par, "paragraphs"):
        par = cell_or_par.add_paragraph() if new_par else cell_or_par.paragraphs[0]
    else:
        par = cell_or_par
    par.paragraph_format.space_after = Pt(0)
    par.paragraph_format.space_before = Pt(0)
    run = par.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    return par


def para(doc, text="", size=10.5, bold=False, italic=False, color=INK, space_after=6, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    if align:
        p.alignment = align
    if text:
        r = p.add_run(text)
        r.font.size = Pt(size)
        r.bold = bold
        r.italic = italic
        r.font.color.rgb = color
    return p


def bullet(doc, text, style="List Bullet"):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    r.font.size = Pt(10.5)
    return p


def numbered(doc, items):
    """Numbered list with literal numbers, so separate lists never continue each other's count."""
    for n, text in enumerate(items, 1):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_after = Pt(3)
        pf.left_indent = Cm(0.7)
        pf.first_line_indent = Cm(-0.7)
        pf.tab_stops.add_tab_stop(Cm(0.7))
        r = p.add_run(f"{n}.\t{text}")
        r.font.size = Pt(10.5)


def heading(doc, text, size=16, color=INK, space_before=0, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(size)
    r.font.color.rgb = color
    return p


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_page_number_footer(section, label):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run(f"{label}  ·  Page ")
    r.font.size = Pt(8)
    r.font.color.rgb = MUTED
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    rr = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "16")
    rpr.append(sz)
    rr.append(rpr)
    t = OxmlElement("w:t")
    t.text = "1"
    rr.append(t)
    fld.append(rr)
    p._p.append(fld)


def length_hint(item):
    n = item["chars"]
    count = f"{n} char" if n == 1 else f"{n} chars"
    if not re.search(r"[A-Za-z]", item["text"]):  # figures like "120+" or "9"
        return f"{count} · keep to a similar short figure"
    if n <= 30:
        words = len(item["text"].split())
        lo, hi = max(1, words - 1), words + 1
        return f"{count} · keep to {lo}–{hi} words"
    return f"{count} · keep near this length"


def simple_table(doc, headers, widths, rows, header_fill=HEADER_FILL, size=9.5):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    fixed_layout(t, widths)
    set_borders(t)
    set_cell_margins(t)
    for i, (h, w) in enumerate(zip(headers, widths)):
        c = t.rows[0].cells[i]
        c.width = Cm(w)
        shade(c, header_fill)
        write(c, h, size=size, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    repeat_header(t.rows[0])
    for row in rows:
        cells = t.add_row().cells
        for i, (v, w) in enumerate(zip(row, widths)):
            cells[i].width = Cm(w)
            bold = isinstance(v, tuple)
            write(cells[i], v[0] if bold else str(v), size=size, bold=bold)
    return t


# ---------- document sections ----------

def cover(doc, meta, pages, items):
    proj = meta.get("project", "Website")
    para(doc, "", space_after=40)
    para(doc, proj.upper(), size=12, bold=True, color=ACCENT, space_after=2)
    para(doc, meta.get("title", "Website Copy Review"), size=30, bold=True, space_after=4)
    sub = " · ".join(x for x in [meta.get("site_url"), meta.get("date") or date.today().strftime("%d %B %Y"),
                                 f"Prepared by {meta['prepared_by']}" if meta.get("prepared_by") else None] if x)
    para(doc, sub, size=11, color=MUTED, space_after=18)
    para(doc, f"{len(items)} items of website wording across {len(pages)} pages and site-wide blocks, "
              f"set out page by page for you to read, check and mark up.", size=11.5, space_after=14)

    scope = meta.get("scope", {})
    heading(doc, "What this document covers", size=13, space_before=4)
    if scope.get("copy"):
        para(doc, scope["copy"])
    if scope.get("content_state"):
        para(doc, scope["content_state"])
    if meta.get("excluded"):
        para(doc, "Not included:", bold=True, space_after=2)
        for x in meta["excluded"]:
            bullet(doc, x)

    heading(doc, "How to review", size=13, space_before=10)
    steps = [
        "Each page section starts with its address on the site. Open the live page alongside this document so you can see the words in place.",
        scope.get("markup") or "Write any new wording in the blank “Suggested revision” column. Leave the row blank if you are happy with the current copy.",
        "Use the “Notes” column for questions or comments that are not replacement wording.",
        "Rows marked with a character count sit in fixed spaces on the page (headings, buttons, menu items). Keep your revision close to the stated length, or the layout may break.",
        "Rows shaded yellow contain an automatic value written like {{this}}. Keep that part exactly as it is.",
        "Please do not renumber, reorder or delete Ref codes or rows. The Ref code is how each change finds its way back to the right place on the site.",
        "Check facts as carefully as wording: founding years, staff numbers, figures, partner and product names, people’s names and job titles. These drift quietly as a site changes.",
    ]
    numbered(doc, steps)


def contents(doc, meta, pages, items):
    page_break(doc)
    heading(doc, "Contents", size=18, space_after=8)
    rows = []
    for p in pages:
        n = sum(1 for i in items if i["page_key"] == p["key"])
        rows.append([p["key"], p["title"], p.get("url", ""), p.get("source", ""), str(n)])
    rows.append([("",), ("Total",), ("",), ("",), (str(len(items)),)])
    simple_table(doc, ["Ref prefix", "Page", "Address on site", "Source file", "Items"],
                 [3.0, 7.0, 6.5, 8.0, 2.4], rows)
    empties = meta.get("empty_templates") or []
    if empties:
        para(doc, "", space_after=4)
        heading(doc, "Templates checked that contain no hardcoded wording", size=11.5, space_before=6)
        for e in empties:
            if isinstance(e, dict):
                bullet(doc, f"{e.get('file')} — {e.get('reason', '')}".rstrip(" —"))
            else:
                bullet(doc, str(e))


def page_section(doc, page, page_items):
    page_break(doc)
    heading(doc, page["title"], size=18, space_after=2)
    meta_bits = [f"Address: {page.get('url', '—')}", f"Source: {page.get('source', '—')}",
                 f"Ref prefix: {page['key']}", f"{len(page_items)} items"]
    para(doc, "   ·   ".join(meta_bits), size=9.5, color=MUTED, space_after=4)
    if page.get("note"):
        para(doc, page["note"], size=10, italic=True, color=MUTED, space_after=6)
    if not page_items:
        para(doc, "No hardcoded wording found for this page.", italic=True)
        return

    widths = [w for _, w in COLS]
    t = doc.add_table(rows=1, cols=len(COLS))
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    fixed_layout(t, widths)
    set_borders(t)
    set_cell_margins(t)
    hdr = t.rows[0]
    for i, (name, w) in enumerate(COLS):
        c = hdr.cells[i]
        c.width = Cm(w)
        shade(c, HEADER_FILL)
        write(c, name, size=9, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    repeat_header(hdr)
    cant_split(hdr)

    current_section = None
    zebra = False
    for it in page_items:
        if it["section"] != current_section:
            current_section = it["section"]
            row = t.add_row()
            cant_split(row)
            merged = row.cells[0].merge(row.cells[-1])
            shade(merged, BANNER_FILL)
            write(merged, current_section.upper(), size=8.5, bold=True, color=ACCENT)
            zebra = False
        row = t.add_row()
        if it["chars"] < 600:  # keep short rows whole; very long ones must be allowed to flow
            cant_split(row)
        cells = row.cells
        for c, w in zip(cells, widths):
            c.width = Cm(w)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            if zebra:
                shade(c, ZEBRA_FILL)
        zebra = not zebra
        write(cells[0], it["ref"], size=8.5, bold=True)
        write(cells[1], it["element"], size=8.5)
        if it.get("constrained"):
            write(cells[1], length_hint(it), size=7.5, italic=True, color=MUTED, new_par=True)
        if it.get("has_dynamic"):
            write(cells[1], "Contains an automatic value — keep the {{…}} part", size=7.5, italic=True,
                  color=RGBColor(0x8A, 0x5A, 0x00), new_par=True)
            shade(cells[2], DYNAMIC_FILL)
        write(cells[2], it["text"], size=9.5)
        # Suggested revision (3) and Notes (4) intentionally left empty.


def closing(doc, meta):
    page_break(doc)
    heading(doc, "What happens next", size=18, space_after=8)
    numbered(doc, [
        "Return this document with your suggested revisions and notes filled in. Please keep the file in Word format.",
        "We apply each revision to the website using its Ref code, so every change lands in the right place.",
        "If a revision does not fit the space it has to go in (for example a long heading or button label), we will come back to you with options rather than shortening it ourselves.",
        "Factual corrections (dates, figures, names, titles) are applied as given. Please flag any you are unsure of in the Notes column.",
        "We will share the updated pages for a final check before they go live.",
    ])
    heading(doc, "Reviewer sign-off", size=13, space_before=14, space_after=6)
    t = simple_table(doc, ["Name", "Role", "Date reviewed"], [9.0, 9.0, 6.0], [["", "", ""]] * 3, size=10)
    for row in t.rows[1:]:
        row.height = Cm(1.0)


# ---------- self-check ----------

def self_check(path, n_items):
    doc = Document(path)
    data_rows = 0
    errors = []
    for t in doc.tables:
        if len(t.columns) != len(COLS) or t.rows[0].cells[0].text.strip() != "Ref":
            continue
        for r in t.rows[1:]:
            cells = r.cells
            ref = cells[0].text.strip()
            if not REF_RE.match(ref):
                continue  # banner row
            data_rows += 1
            if cells[3].text.strip() or cells[4].text.strip():
                errors.append(f"{ref}: revision/notes cell is not empty")
            if not cells[2].text.strip():
                errors.append(f"{ref}: current copy is empty")
    if data_rows != n_items:
        errors.append(f"data rows {data_rows} != items {n_items}")
    return data_rows, errors


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--pages", required=True)
    ap.add_argument("--items", required=True)
    ap.add_argument("--out", required=True)
    args = ap.parse_args()

    cfg = json.load(open(args.pages))
    meta, pages = cfg.get("meta", {}), cfg["pages"]
    items = json.load(open(args.items))
    missing = [i for i in items if "ref" not in i]
    if missing:
        sys.exit("items.json has entries without 'ref' — run merge_and_verify.py first")

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(1.4))
    sec.footer_distance = Cm(0.7)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    add_page_number_footer(sec, f"{meta.get('project', 'Website')} — {meta.get('title', 'Website Copy Review')}")

    cover(doc, meta, pages, items)
    contents(doc, meta, pages, items)
    for p in pages:
        page_section(doc, p, [i for i in items if i["page_key"] == p["key"]])
    closing(doc, meta)
    doc.save(args.out)

    rows, errors = self_check(args.out, len(items))
    print(f"Wrote {args.out}: {rows} data rows across {len(pages)} page sections")
    if errors:
        print("SELF-CHECK FAILED:")
        for e in errors:
            print("  -", e)
        sys.exit(1)
    print("Self-check passed: row count matches, revision and notes columns empty.")


if __name__ == "__main__":
    main()
